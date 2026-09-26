# Conversor de Excel para CSV

import openpyxl, csv, os
from pathlib import Path

caminho = (Path.home() / "Documents" / "VSCode" / "Automate_the_Boring_Stuff_3e_onlinematerials")

for excel_file in os.listdir(caminho):
    
    # Skip non-xlsx files, load the workbook object
    if not excel_file.endswith('.xlsx'):
        continue

    try:

        wb = openpyxl.load_workbook(caminho / excel_file)

    except Exception as error:
        print(f'Não foi possível acessar o {excel_file} devido ao erro: {error}')
        continue

    for sheet_name in wb.sheetnames:

        # Loop through every sheet in the workbook.
        sheet = wb[sheet_name]

        # Create the CSV filename from the Excel filename and sheet title.
        csv_name_archive = (Path(excel_file).stem) + '_' + sheet_name + '.csv'

        # Create the csv.writer object for this CSV file.
        try:

            csv_file = open(csv_name_archive, 'w', newline='')
            csv_writer = csv.writer(csv_file)
        
            # Loop through every row in the sheet.
            for row_num in range(1, sheet.max_row + 1):
                row_data = []    # Append each cell to this list.

                # Loop through each cell in the row.
                for col_num in range(1, sheet.max_column + 1):

                    # Append each cell's data to row_data
                    row_data.append(str(sheet.cell(row =row_num, column = col_num).value))

                # Write the row_data list to the CSV file.
                csv_writer.writerow(row_data)

            csv_file.close()

        except Exception as csv_error:
                    print(
                f"Erro ao converter {excel_file} "
                f"(planilha {sheet_name}): {csv_error}")
                    continue 
